"""Render docs/SOLUTION.md (+ appendices: full prompt, full filter list) to docs/WinWin_AI_Assistant.docx.

The .docx is meant for upload to Google Drive -> "Open with Google Docs".
Run:  pip install python-docx && python docs/build_docx.py
"""

from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "WinWin_AI_Assistant.docx"
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_inline(par, text: str) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2].replace("`", "")).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def add_code(doc, lines: list[str]) -> None:
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(6)
    run = par.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(8)


def add_table(doc, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), "D9D9D9")
        borders.append(tag)
    table._tbl.tblPr.append(borders)
    for r, cells in enumerate(rows):
        for c, text in enumerate(cells):
            cell = table.cell(r, c)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            add_inline(cell.paragraphs[0], text.strip())
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(8.5)
                if r == 0:
                    run.bold = True
            if r == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "ECECEC")
                cell._tc.get_or_add_tcPr().append(shade)
        if r == 0:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            table.rows[r]._tr.get_or_add_trPr().append(header)
    doc.add_paragraph()


def render_markdown(doc, md: str, heading_offset: int = 0) -> None:
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            add_code(doc, block)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[\s:\-|]+\|$", lines[i]):
                    rows.append(lines[i].strip("|").split("|"))
                i += 1
            add_table(doc, rows)
            continue
        elif m := re.match(r"^(#{1,4})\s+(.*)", line):
            level = min(len(m.group(1)) + heading_offset, 4)
            if level == 1:
                h = doc.add_heading(m.group(2), 0)
                h.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                doc.add_heading(m.group(2), level - 1)
        elif m := re.match(r"^(\s*)[*-]\s+(.*)", line):
            text = m.group(2)
            while i + 1 < len(lines) and re.match(r"^\s{2,}\S", lines[i + 1]) and not re.match(r"^\s*[*-]\s", lines[i + 1]):
                i += 1
                text += " " + lines[i].strip()
            add_inline(doc.add_paragraph(style="List Bullet 2" if m.group(1) else "List Bullet"), text)
        elif m := re.match(r"^(\d+)\.\s+(.*)", line):
            # manual numbers: Word's "List Number" style would continue numbering across sections
            text = f"{m.group(1)}. {m.group(2)}"
            while i + 1 < len(lines) and re.match(r"^\s{2,}\S", lines[i + 1]):
                i += 1
                text += " " + lines[i].strip()
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Pt(18)
            par.paragraph_format.first_line_indent = Pt(-12)
            add_inline(par, text)
        elif line.strip() == "---":
            pass
        elif line.strip():
            text = line
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"^(#|\||```|\s*[*-]\s|\d+\.\s|---)", lines[i + 1]):
                i += 1
                text += " " + lines[i].strip()
            add_inline(doc.add_paragraph(), text)
        i += 1


def main() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.7)
    sec.left_margin = sec.right_margin = Inches(.72)
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(5)
    for name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12),
                       ("Heading 3", 10.5), ("Heading 4", 10)):
        heading = doc.styles[name]
        heading.font.name = "Arial"
        heading.font.size = Pt(size)
        heading.font.color.rgb = RGBColor(0, 0, 0)
        heading.paragraph_format.keep_with_next = True

    render_markdown(doc, (ROOT / "docs" / "SOLUTION.md").read_text(encoding="utf-8"))

    doc.add_page_break()
    doc.add_heading("Appendix A Full system prompt", 1)
    render_markdown(doc, (ROOT / "assistant" / "system_prompt.md").read_text(encoding="utf-8"), heading_offset=1)

    doc.add_page_break()
    catalog = json.loads((ROOT / "filters" / "catalog.json").read_text(encoding="utf-8"))
    doc.add_heading(f"Appendix B Full filter list {catalog['count']} filters", 1)
    add_inline(doc.add_paragraph(), "Also available as `filters/catalog.csv` with synonyms and conflict metadata. "
                                    "Range filters show unit and bounds; ⚔ = has hard conflicts, ⇒ = implies other filters.")
    by_cat: dict[str, list[dict]] = {}
    for f in catalog["filters"]:
        by_cat.setdefault(f["category"], []).append(f)
    for cat, items in by_cat.items():
        doc.add_heading(f"{cat} ({len(items)})", 3)
        rows = [["id", "label", "type"]]
        for f in items:
            t = f["type"]
            if t == "range":
                t = f"range {f['min']}–{f['max']} {f['unit']}"
            elif t == "enum":
                t = "enum: " + ", ".join(f["options"][:6]) + ("…" if len(f["options"]) > 6 else "")
            marks = (" ⚔" if f.get("conflicts_with") or f.get("exclusive_group") else "") + (" ⇒" if f.get("implies") else "")
            rows.append([f"`{f['id']}`", f["label"] + marks, t])
        add_table(doc, rows)

    doc.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
